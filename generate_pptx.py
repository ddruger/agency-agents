from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Colors ──────────────────────────────────────────────────────────────────
DARK       = RGBColor(13,  13,  13)
DARK2      = RGBColor(26,  26,  26)
DARKEST    = RGBColor(10,  10,  10)
EMBER      = RGBColor(201, 75,  12)
EMBER_B    = RGBColor(232, 93,  4)
EMBER_GOLD = RGBColor(244, 140, 6)
EMBER_LT   = RGBColor(255, 107, 53)
WHITE      = RGBColor(255, 255, 255)
SECONDARY  = RGBColor(179, 179, 179)
CARD_BG    = RGBColor(26,  26,  26)

# ── Presentation setup ──────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

blank_layout = prs.slide_layouts[6]  # completely blank

# ── Helper functions ─────────────────────────────────────────────────────────

def add_bg(slide, color):
    """Full-slide background rectangle."""
    shape = slide.shapes.add_shape(
        1, 0, 0, SLIDE_W, SLIDE_H
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_rect(slide, left, top, width, height, fill_color, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def add_tb(slide, left, top, width, height, text, font_name, font_size, bold, color,
           align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txb


def add_label(slide, text, top=Inches(0.5)):
    add_tb(slide, Inches(0.6), top, Inches(4), Inches(0.4),
           text, "Arial", 10, True, EMBER, PP_ALIGN.LEFT)


def add_section_title(slide, line1, line2, top=Inches(1.1), size=64):
    add_tb(slide, Inches(0.6), top, Inches(8), Inches(1),
           line1, "Arial Black", size, True, WHITE, PP_ALIGN.LEFT)
    if line2:
        add_tb(slide, Inches(0.6), top + Inches(0.95), Inches(8), Inches(1),
               line2, "Arial Black", size, True, WHITE, PP_ALIGN.LEFT)


def multi_para_tb(slide, left, top, width, height, paras, font_name="Arial",
                  font_size=12, bold=False, color=WHITE, align=PP_ALIGN.LEFT, spacing=None):
    """Add a textbox with multiple paragraphs."""
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    for i, (text, fname, fsize, fbold, fcolor, falign) in enumerate(paras):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = falign
        run = p.add_run()
        run.text = text
        run.font.name = fname
        run.font.size = Pt(fsize)
        run.font.bold = fbold
        run.font.color.rgb = fcolor
        if spacing and i > 0:
            p.space_before = Pt(spacing)
    return txb


def pill_label(slide, left, top, text):
    """Small ember-outlined pill badge."""
    w, h = Inches(1.4), Inches(0.32)
    r = add_rect(slide, left, top, w, h, DARK2, EMBER, Pt(1))
    add_tb(slide, left, top, w, h, text, "Arial", 8, True, EMBER, PP_ALIGN.CENTER)


def card(slide, left, top, width, height, title, body, title_size=13):
    add_rect(slide, left, top, width, height, DARK2, EMBER, Pt(1.5))
    add_tb(slide, left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), Inches(0.35),
           title, "Arial Black", title_size, True, EMBER, PP_ALIGN.LEFT)
    add_tb(slide, left + Inches(0.15), top + Inches(0.52), width - Inches(0.3), height - Inches(0.65),
           body, "Arial", 10, False, SECONDARY, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title / Hero
# ════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank_layout)
add_bg(s1, DARK)

# Ember vertical accent bar on right
add_rect(s1, SLIDE_W - Inches(0.18), 0, Inches(0.1), SLIDE_H, EMBER_B)

# Thin horizontal ember line near top
add_rect(s1, Inches(0.6), Inches(0.52), Inches(0.83), Inches(0.05), EMBER)

# Label
add_tb(s1, Inches(0.6), Inches(0.65), Inches(5), Inches(0.35),
       "THE BUSINESS PODCAST", "Arial", 9, True, EMBER, PP_ALIGN.LEFT)

# Main title "NO"
add_tb(s1, Inches(0.55), Inches(1.15), Inches(10), Inches(1.8),
       "NO", "Arial Black", 88, True, WHITE, PP_ALIGN.LEFT)

# Main title "PERMISSION"
add_tb(s1, Inches(0.55), Inches(2.7), Inches(10), Inches(1.8),
       "PERMISSION", "Arial Black", 88, True, WHITE, PP_ALIGN.LEFT)

# Tagline
add_tb(s1, Inches(0.6), Inches(4.55), Inches(8), Inches(0.5),
       "Built for Founders Who Don't Wait", "Arial", 20, False, SECONDARY, PP_ALIGN.LEFT)

# Bottom separator line
add_rect(s1, Inches(0.6), Inches(6.8), Inches(12.1), Inches(0.04), EMBER)

# Website
add_tb(s1, Inches(0.6), Inches(6.9), Inches(5), Inches(0.4),
       "nopermissionpod.com", "Arial", 11, False, SECONDARY, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — About the Show
# ════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank_layout)
add_bg(s2, RGBColor(17, 17, 17))

add_label(s2, "THE SHOW")
add_section_title(s2, "Built Different.", "Built Bold.", top=Inches(0.9), size=60)

# Three pillar boxes
pillars = [
    ("REAL TALK",     "Unfiltered conversations with people who've done the work"),
    ("NO FILTERS",    "Raw insights, hard lessons, and the specifics that matter"),
    ("NO PERMISSION", "The show built for those who don't wait for validation"),
]
col_w   = Inches(3.8)
col_gap = Inches(0.2)
col_top = Inches(3.0)
col_h   = Inches(1.85)
for i, (t, b) in enumerate(pillars):
    lx = Inches(0.6) + i * (col_w + col_gap)
    card(s2, lx, col_top, col_w, col_h, t, b, title_size=13)

# Show description
add_tb(s2, Inches(0.6), Inches(5.05), Inches(12.0), Inches(1.5),
       "No Permission is the weekly business podcast for founders, operators, and investors "
       "who build on their own terms. No theory. No fluff. Just the real playbook.",
       "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Meet Your Host
# ════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank_layout)
add_bg(s3, DARK)

add_label(s3, "MEET YOUR HOST")

# Avatar circle (ellipse)
from pptx.util import Inches as IN
from pptx.enum.dml import MSO_THEME_COLOR
avatar_l, avatar_t = Inches(0.6), Inches(0.85)
avatar_d = Inches(2.2)
av = s3.shapes.add_shape(9, avatar_l, avatar_t, avatar_d, avatar_d)  # 9 = oval
av.fill.solid()
av.fill.fore_color.rgb = EMBER
av.line.fill.background()
add_tb(s3, avatar_l, avatar_t + Inches(0.55), avatar_d, avatar_d - Inches(0.5),
       "DD", "Arial Black", 44, True, WHITE, PP_ALIGN.CENTER)

# Name & title
add_tb(s3, Inches(0.6), Inches(3.2), Inches(3.5), Inches(0.6),
       "Daniel Druger", "Arial Black", 28, True, WHITE, PP_ALIGN.LEFT)
add_tb(s3, Inches(0.6), Inches(3.75), Inches(3.5), Inches(0.4),
       "Host · No Permission", "Arial", 12, False, EMBER, PP_ALIGN.LEFT)

# Badge pills
badges = ["Co-Host, ADSN", "Active Investor", "Startup Advisor"]
for bi, bt in enumerate(badges):
    pill_label(s3, Inches(0.6) + bi * Inches(1.55), Inches(4.28), bt)

# Right side bio
right_l = Inches(4.2)
add_tb(s3, right_l, Inches(0.85), Inches(8.7), Inches(0.9),
       "Built From the Inside Out", "Arial Black", 40, True, WHITE, PP_ALIGN.LEFT)

bio1 = ("Daniel Druger isn't talking about business from the sidelines. As co-host of ADSN, "
        "an active investor, and hands-on advisor to founders, Daniel brings real capital, "
        "real relationships, and real-world decisions to every conversation.")
add_tb(s3, right_l, Inches(1.9), Inches(8.7), Inches(1.5),
       bio1, "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)

bio2 = ("No Permission exists because the best business lessons aren't found in books — "
        "they're found in the rooms most people can't get into. Daniel opens that door.")
add_tb(s3, right_l, Inches(3.5), Inches(8.7), Inches(1.2),
       bio2, "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)

# Tag pills on right
tags = ["INVESTOR", "ADVISOR", "OPERATOR", "ADSN CO-HOST"]
for ti, tt in enumerate(tags):
    pill_label(s3, right_l + ti * Inches(1.55), Inches(4.9), tt)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Who It's For
# ════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(blank_layout)
add_bg(s4, RGBColor(17, 17, 17))

add_label(s4, "THE AUDIENCE")
add_section_title(s4, "Built for Founders", "Who Don't Wait", top=Inches(0.9), size=52)

audience = [
    ("ENTREPRENEURS", "Founders at every stage building companies from scratch or scaling to the next level."),
    ("OPERATORS",     "The executives and managers who execute strategy and turn vision into reality every day."),
    ("INVESTORS",     "Angels, VCs, and family offices looking for pattern recognition from proven operators."),
    ("BUILDERS",      "Engineers, marketers, and creatives turning ideas into products and careers into legacies."),
]
card_w, card_h = Inches(5.9), Inches(1.55)
gap = Inches(0.25)
for i, (t, b) in enumerate(audience):
    row, col = divmod(i, 2)
    lx = Inches(0.6) + col * (card_w + gap)
    ly = Inches(3.2) + row * (card_h + gap)
    card(s4, lx, ly, card_w, card_h, t, b)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — What We Cover
# ════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(blank_layout)
add_bg(s5, DARK)

add_label(s5, "TOPICS")
add_tb(s5, Inches(0.6), Inches(0.85), Inches(10), Inches(1.0),
       "What We Cover", "Arial Black", 60, True, WHITE, PP_ALIGN.LEFT)

topics = [
    ("BUILDING COMPANIES",  "Zero to one, product-market fit, hiring your first team"),
    ("RAISING CAPITAL",     "Fundraising strategy, investor relations, term sheets"),
    ("MARKETING & GROWTH",  "Acquisition, retention, brand building at scale"),
    ("LEADERSHIP",          "Managing up, down, and across in high-growth environments"),
    ("SALES & REVENUE",     "Pipeline, deals, pricing, and revenue operations"),
    ("CULTURE & TEAM",      "Hiring, retention, and values in action"),
]
card_w, card_h = Inches(3.9), Inches(1.45)
gap_x, gap_y   = Inches(0.2), Inches(0.22)
for i, (t, b) in enumerate(topics):
    row, col = divmod(i, 3)
    lx = Inches(0.6) + col * (card_w + gap_x)
    ly = Inches(2.1) + row * (card_h + gap_y)
    card(s5, lx, ly, card_w, card_h, t, b)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — The Format
# ════════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(blank_layout)
add_bg(s6, RGBColor(17, 17, 17))

add_label(s6, "THE FORMAT")
add_section_title(s6, "How We", "Deliver It", top=Inches(0.9), size=60)

add_tb(s6, Inches(0.6), Inches(3.05), Inches(6.2), Inches(1.0),
       "No Permission episodes are crafted to respect your time and maximize your ROI per minute listened.",
       "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)

# Stats row
stats = [("WEEKLY", "Release cadence"), ("45-60 MIN", "Per episode"), ("SINCE 2024", "In your feed")]
stat_w = Inches(1.9)
for si, (sv, sl) in enumerate(stats):
    sx = Inches(0.6) + si * Inches(2.1)
    add_tb(s6, sx, Inches(4.2), stat_w, Inches(0.45), sv, "Arial Black", 18, True, EMBER, PP_ALIGN.LEFT)
    add_tb(s6, sx, Inches(4.65), stat_w, Inches(0.3), sl, "Arial", 10, False, SECONDARY, PP_ALIGN.LEFT)

# Right side format cards
formats = [
    ("SOLO DEEP-DIVES",  "Host-driven tactical breakdowns on a single business topic."),
    ("GUEST INTERVIEWS", "Long-form conversations with founders, operators, and investors."),
    ("TACTICAL Q&A",     "Direct listener questions answered with no fluff."),
]
fmt_w, fmt_h = Inches(5.9), Inches(1.4)
fmt_x = Inches(7.0)
for fi, (t, b) in enumerate(formats):
    fy = Inches(1.0) + fi * (fmt_h + Inches(0.2))
    card(s6, fmt_x, fy, fmt_w, fmt_h, t, b)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Why Listen / The Numbers
# ════════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(blank_layout)
add_bg(s7, DARK)

add_label(s7, "WHY NO PERMISSION")
add_section_title(s7, "The Numbers", "Speak First", top=Inches(0.9), size=60)

stat_boxes = [
    ("100%",    "Guest Vetting Rate — every guest has an operational track record"),
    ("ZERO",    "Paid Guest Placements — merit only, always"),
    ("45-60 MIN","Per Episode — deep enough to matter"),
    ("WEEKLY",  "New Episodes — consistent, reliable, always on"),
]
sb_w, sb_h = Inches(2.9), Inches(1.8)
sb_gap = Inches(0.25)
for si, (sv, sl) in enumerate(stat_boxes):
    sx = Inches(0.6) + si * (sb_w + sb_gap)
    sy = Inches(3.0)
    add_rect(s7, sx, sy, sb_w, sb_h, DARK2, EMBER, Pt(1.5))
    add_tb(s7, sx + Inches(0.15), sy + Inches(0.15), sb_w - Inches(0.3), Inches(0.65),
           sv, "Arial Black", 22, True, EMBER_B, PP_ALIGN.LEFT)
    add_tb(s7, sx + Inches(0.15), sy + Inches(0.85), sb_w - Inches(0.3), sb_h - Inches(1.0),
           sl, "Arial", 11, False, SECONDARY, PP_ALIGN.LEFT)

add_tb(s7, Inches(0.6), Inches(5.15), Inches(12), Inches(0.45),
       "No fluff. No gatekeeping. No Permission.", "Arial Black", 20, True, WHITE, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Guest Spotlight
# ════════════════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(blank_layout)
add_bg(s8, RGBColor(17, 17, 17))

add_label(s8, "THIS WEEK'S GUEST")

# Left card background
add_rect(s8, Inches(0.5), Inches(0.85), Inches(5.3), Inches(6.1), DARK2, EMBER, Pt(1.5))

# Avatar circle
av2 = s8.shapes.add_shape(9, Inches(0.85), Inches(1.0), Inches(1.8), Inches(1.8))
av2.fill.solid(); av2.fill.fore_color.rgb = EMBER; av2.line.fill.background()
add_tb(s8, Inches(0.85), Inches(1.35), Inches(1.8), Inches(1.0),
       "G", "Arial Black", 44, True, WHITE, PP_ALIGN.CENTER)

# Guest info
add_tb(s8, Inches(0.75), Inches(2.95), Inches(4.8), Inches(0.55),
       "Guest Name", "Arial Black", 28, True, WHITE, PP_ALIGN.LEFT)
add_tb(s8, Inches(0.75), Inches(3.5), Inches(4.8), Inches(0.4),
       "Title · Company Name", "Arial", 13, False, EMBER, PP_ALIGN.LEFT)

add_rect(s8, Inches(0.75), Inches(4.0), Inches(2.5), Inches(0.3), EMBER)
add_tb(s8, Inches(0.75), Inches(4.0), Inches(4.8), Inches(0.3),
       "EPISODE TOPIC", "Arial", 9, True, WHITE, PP_ALIGN.LEFT)

add_tb(s8, Inches(0.75), Inches(4.4), Inches(4.8), Inches(1.0),
       "The tactical framework this guest used to scale from $0 to $10M ARR in under two years.",
       "Arial", 11, False, SECONDARY, PP_ALIGN.LEFT)

add_tb(s8, Inches(0.75), Inches(5.5), Inches(4.8), Inches(0.4),
       "Episode 101 · 52 min", "Arial", 11, False, EMBER_GOLD, PP_ALIGN.LEFT)

# Right side
right_x = Inches(6.2)
add_tb(s8, right_x, Inches(0.5), Inches(6.8), Inches(0.35),
       "GUEST SPOTLIGHT", "Arial", 10, True, EMBER, PP_ALIGN.LEFT)
add_tb(s8, right_x, Inches(0.95), Inches(6.8), Inches(0.75),
       "Real People.", "Arial Black", 44, True, WHITE, PP_ALIGN.LEFT)
add_tb(s8, right_x, Inches(1.65), Inches(6.8), Inches(0.75),
       "Real Receipts.", "Arial Black", 44, True, WHITE, PP_ALIGN.LEFT)

desc = ("Every guest on No Permission has done the work. "
        "We vet for operational experience, not follower count.")
add_tb(s8, right_x, Inches(2.55), Inches(6.8), Inches(0.8),
       desc, "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)

bullets = [
    "Verified track record of building or operating",
    "Specific metrics, decisions, and outcomes",
    "No PR fluff — real conversation, real answers",
    "Curated for founders who need the real playbook",
]
for bi, bt in enumerate(bullets):
    add_tb(s8, right_x, Inches(3.55) + bi * Inches(0.6), Inches(6.8), Inches(0.5),
           f"◆  {bt}", "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — Connect & Subscribe
# ════════════════════════════════════════════════════════════════════════════
s9 = prs.slides.add_slide(blank_layout)
add_bg(s9, DARK)

add_label(s9, "LISTEN EVERYWHERE")
add_section_title(s9, "Find Us On", "Every Platform", top=Inches(0.9), size=56)

platforms = ["SPOTIFY", "APPLE PODCASTS", "YOUTUBE"]
plat_w, plat_h = Inches(3.8), Inches(1.5)
plat_gap = Inches(0.25)
for pi, pn in enumerate(platforms):
    px = Inches(0.6) + pi * (plat_w + plat_gap)
    add_rect(s9, px, Inches(3.0), plat_w, plat_h, DARK2, EMBER, Pt(1.5))
    add_tb(s9, px + Inches(0.15), Inches(3.0) + Inches(0.55), plat_w - Inches(0.3), Inches(0.55),
           pn, "Arial Black", 18, True, WHITE, PP_ALIGN.CENTER)

# Social row
socials = [
    ("Instagram",  "@nopermissionpod"),
    ("X / Twitter","@nopermissionpod"),
    ("LinkedIn",   "No Permission Podcast"),
    ("TikTok",     "@nopermissionpod"),
]
add_rect(s9, Inches(0.6), Inches(4.75), Inches(12.1), Inches(0.05), EMBER)
soc_w = Inches(2.9)
for si, (sn, sh) in enumerate(socials):
    sx = Inches(0.6) + si * (soc_w + Inches(0.1))
    add_tb(s9, sx, Inches(4.9), soc_w, Inches(0.35), sn, "Arial", 10, True, EMBER, PP_ALIGN.LEFT)
    add_tb(s9, sx, Inches(5.25), soc_w, Inches(0.35), sh, "Arial", 11, False, SECONDARY, PP_ALIGN.LEFT)

# CTA
add_tb(s9, Inches(0.6), Inches(5.85), Inches(12), Inches(0.5),
       "Start Listening. No Permission Required.", "Arial Black", 22, True, WHITE, PP_ALIGN.CENTER)
add_tb(s9, Inches(0.6), Inches(6.4), Inches(12), Inches(0.4),
       "nopermissionpod.com", "Arial", 14, False, EMBER, PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — Sponsorship
# ════════════════════════════════════════════════════════════════════════════
s10 = prs.slides.add_slide(blank_layout)
add_bg(s10, RGBColor(17, 17, 17))

add_label(s10, "PARTNERSHIPS")
add_section_title(s10, "Partner", "With Us", top=Inches(0.9), size=60)

# Left description
add_tb(s10, Inches(0.6), Inches(3.05), Inches(5.8), Inches(0.8),
       "Reach an audience of decision-makers, founders, and operators who act on what they hear.",
       "Arial", 12, False, SECONDARY, PP_ALIGN.LEFT)

demo_stats = [
    ("78%",   "Decision Makers"),
    ("34",    "Avg. Audience Age"),
    ("$120K+","Household Income"),
    ("65%",   "Own a Business"),
]
ds_w = Inches(1.3)
for di, (dv, dl) in enumerate(demo_stats):
    dx = Inches(0.6) + di * Inches(1.55)
    add_tb(s10, dx, Inches(4.05), ds_w, Inches(0.45), dv, "Arial Black", 18, True, EMBER, PP_ALIGN.LEFT)
    add_tb(s10, dx, Inches(4.5), ds_w, Inches(0.35), dl, "Arial", 9, False, SECONDARY, PP_ALIGN.LEFT)

# Right sponsorship tiers
tiers = [
    ("PRESENTING SPONSOR",
     "Full episode ownership, 60-sec mid-roll + pre-roll, social amplification"),
    ("FEATURED PARTNER",
     "30-sec mid-roll, newsletter mention, episode show notes link"),
    ("BRAND MENTION",
     "Host-read mention, show notes placement, social tag"),
]
tier_w, tier_h = Inches(6.2), Inches(1.35)
tier_x = Inches(6.8)
for ti, (tt, tb_) in enumerate(tiers):
    ty = Inches(1.0) + ti * (tier_h + Inches(0.2))
    card(s10, tier_x, ty, tier_w, tier_h, tt, tb_)

# Contact
add_rect(s10, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.04), EMBER)
add_tb(s10, Inches(0.6), Inches(6.65), Inches(12), Inches(0.4),
       "partnerships@nopermissionpod.com", "Arial", 12, False, EMBER, PP_ALIGN.LEFT)


# ════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Closing
# ════════════════════════════════════════════════════════════════════════════
s11 = prs.slides.add_slide(blank_layout)
add_bg(s11, DARKEST)

# Top ember bar
add_rect(s11, 0, Inches(0.22), SLIDE_W, Inches(0.07), EMBER)
# Bottom ember bar
add_rect(s11, 0, SLIDE_H - Inches(0.3), SLIDE_W, Inches(0.07), EMBER)

# Subtitle label
add_tb(s11, 0, Inches(1.6), SLIDE_W, Inches(0.5),
       "— The Business Podcast —", "Arial", 13, True, EMBER, PP_ALIGN.CENTER)

# Main "NO"
add_tb(s11, 0, Inches(2.0), SLIDE_W, Inches(1.8),
       "NO", "Arial Black", 96, True, WHITE, PP_ALIGN.CENTER)

# Main "PERMISSION"
add_tb(s11, 0, Inches(3.6), SLIDE_W, Inches(1.8),
       "PERMISSION", "Arial Black", 96, True, WHITE, PP_ALIGN.CENTER)

# Tagline
add_tb(s11, 0, Inches(5.4), SLIDE_W, Inches(0.5),
       "New Episodes Every Week", "Arial", 20, False, SECONDARY, PP_ALIGN.CENTER)

# Website
add_tb(s11, 0, Inches(5.95), SLIDE_W, Inches(0.45),
       "nopermissionpod.com", "Arial", 16, False, EMBER, PP_ALIGN.CENTER)

# Platforms
add_tb(s11, 0, Inches(6.45), SLIDE_W, Inches(0.45),
       "Spotify  ◆  Apple Podcasts  ◆  YouTube  ◆  @nopermissionpod",
       "Arial", 11, False, SECONDARY, PP_ALIGN.CENTER)


# ── Save ─────────────────────────────────────────────────────────────────────
output_path = "/home/user/agency-agents/no-permission-deck.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")

import os
size = os.path.getsize(output_path)
print(f"File size: {size:,} bytes ({size/1024:.1f} KB)")
